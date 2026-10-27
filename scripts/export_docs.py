"""Export the design note to Word using a deliberately small Markdown subset.

Convert the resulting DOCX to PDF with LibreOffice, inspect all pages, then
run scripts/record_doc_exports.py. See docs/document-exports.md.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, text):
    pattern = r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))"
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif re.fullmatch(r"\[[^\]]+\]\([^)]+\)", part):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            paragraph.add_run(match.group(1))
        else:
            paragraph.add_run(part)


def table(document, rows):
    grid = document.add_table(rows=1, cols=len(rows[0]))
    grid.autofit = False
    widths = [1.45, 3.55, 1.75] if len(rows[0]) == 3 else [2.05, 4.7]
    for column, width in zip(grid.columns, widths, strict=True):
        column.width = Inches(width)
    for cell, width in zip(grid.rows[0].cells, widths, strict=True):
        cell.width = Inches(width)
    for i, row in enumerate(rows):
        cells = grid.rows[0].cells if i == 0 else grid.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, row, strict=True)):
            cell.width = Inches(widths[j])
            inline(cell.paragraphs[0], text)
            cell.vertical_alignment = 1
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(2)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "172B32")
                    run.bold = i == 0
            props = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "173E49" if i == 0 else ("F1F5F6" if i % 2 else "FFFFFF"))
            props.append(shading)
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right"]:
                node = OxmlElement("w:" + edge)
                for key, value in [("val", "single"), ("sz", "4"), ("color", "D9D9D9")]:
                    node.set(qn("w:" + key), value)
                borders.append(node)
            props.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge in ["top", "left", "bottom", "right"]:
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), "60")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            props.append(margins)
        row_props = grid.rows[i]._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if i == 0:
            row_props.append(OxmlElement("w:tblHeader"))
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.875)
    section.top_margin = section.bottom_margin = Inches(0.75)
    for name, size in [("Normal", 11), ("Title", 24), ("Heading 1", 16), ("Heading 2", 12)]:
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.08
        if name != "Normal":
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(12)
    document.core_properties.title = "Holdline design and operating model"
    document.core_properties.author = "Holdline contributors"
    document.core_properties.subject = "Receiving review architecture and operating decisions"
    document.core_properties.keywords = "Holdline, A2A, MCP, receiving, learning"
    lines = (ROOT / "docs/holdline.md").read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in cells):
                    rows.append(cells)
                i += 1
            table(document, rows)
            continue
        if line.startswith("# "):
            inline(document.add_paragraph(style="Title"), line[2:])
            document.add_paragraph("Version 2.0.0 · 26 October 2026")
        elif line.startswith("## "):
            inline(document.add_paragraph(style="Heading 1"), line[3:])
        elif line.startswith("### "):
            inline(document.add_paragraph(style="Heading 2"), line[4:])
        elif re.match(r"\d+\. ", line):
            inline(document.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        else:
            parts = [line]
            while (
                i + 1 < len(lines)
                and lines[i + 1].strip()
                and not lines[i + 1].startswith(("#", "|", "- "))
            ):
                i += 1
                parts.append(lines[i].strip())
            inline(document.add_paragraph(), " ".join(parts))
        i += 1
    # Remove title/theme paragraph rules from the document template.
    for tree in (document.styles.element, document.element):
        for border in tree.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    destination = ROOT / "docs/holdline.docx"
    document.save(destination)
    print("Word export created:", destination)


if __name__ == "__main__":
    main()
